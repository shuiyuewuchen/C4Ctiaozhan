#!/usr/bin/env python3
"""
Stage 1: 文档摄入（国产模型增强版）
====================================
支持: Markdown / TXT / PDF(文本型) / DOCX
相比 starter kit 新增:
  - PDF 摄入 (pdfplumber)
  - DOCX 摄入 (python-docx)
  - 保留原始 LaTeX 行内公式 ($...$) 和行间公式 ($$...$$)
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path


# ─────────────────────────────────────────────
# 各格式读取器
# ─────────────────────────────────────────────

def read_markdown(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8") as f:
        return f.read()


def read_pdf_text(filepath: str) -> str:
    """文本型 PDF: pdfplumber 逐页提取，保留换行。"""
    try:
        import pdfplumber  # type: ignore
    except ImportError as e:
        raise RuntimeError(
            "PDF 摄入需要 pdfplumber: pip install pdfplumber"
        ) from e

    pages = []
    with pdfplumber.open(filepath) as pdf:
        for p in pdf.pages:
            txt = p.extract_text() or ""
            pages.append(txt)
    return "\n\n".join(pages)


def read_docx(filepath: str) -> str:
    try:
        import docx  # type: ignore
    except ImportError as e:
        raise RuntimeError(
            "DOCX 摄入需要 python-docx: pip install python-docx"
        ) from e

    d = docx.Document(filepath)
    lines = []
    for para in d.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    # 表格也提取出来
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)


# ─────────────────────────────────────────────
# 格式路由
# ─────────────────────────────────────────────

FORMAT_HANDLERS = {
    ".md":   read_markdown,
    ".txt":  read_markdown,
    ".pdf":  read_pdf_text,
    ".docx": read_docx,
}


def detect_format(filepath: str) -> str:
    ext = Path(filepath).suffix.lower()
    if ext not in FORMAT_HANDLERS:
        raise ValueError(
            f"不支持的文件格式: {ext}\n支持: {list(FORMAT_HANDLERS.keys())}"
        )
    return ext


def split_into_sections(text: str) -> list:
    sections = []
    current_title = "Untitled"
    current_lines: list[str] = []

    for line in text.split("\n"):
        heading = re.match(r"^(#{1,4})\s+(.+)", line)
        if heading:
            if current_lines:
                content = "\n".join(current_lines).strip()
                if content:
                    sections.append({"title": current_title, "content": content})
            current_title = heading.group(2).strip()
            current_lines = []
        else:
            current_lines.append(line)

    if current_lines:
        content = "\n".join(current_lines).strip()
        if content:
            sections.append({"title": current_title, "content": content})
    return sections


def ingest(filepath: str) -> dict:
    filepath = str(filepath)
    ext = detect_format(filepath)
    raw_text = FORMAT_HANDLERS[ext](filepath)
    return {
        "source_file": Path(filepath).name,
        "format": ext,
        "raw_text": raw_text,
        "sections": split_into_sections(raw_text),
    }


# ─────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────

def main():
    if len(sys.argv) < 3:
        print("用法: python ingest.py <输入文件> <输出.json>")
        sys.exit(1)
    input_path, output_path = sys.argv[1], sys.argv[2]
    if not Path(input_path).exists():
        print(f"错误: 文件不存在 — {input_path}")
        sys.exit(1)

    print(f"[Stage 1] 文档摄入: {input_path}")
    result = ingest(input_path)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(f"  格式: {result['format']}")
    print(f"  分段: {len(result['sections'])} sections")
    print(f"  输出: {output_path}")


if __name__ == "__main__":
    main()
